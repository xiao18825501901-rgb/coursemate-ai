import re
from collections.abc import Callable
from pathlib import Path

from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader

from app.rag.errors import DocumentLoadError
from app.rag.types import SourceSection

Loader = Callable[[Path], list[SourceSection]]


def _require_content(sections: list[SourceSection], path: Path) -> list[SourceSection]:
    populated = [section for section in sections if section.text.strip()]
    if not populated:
        raise DocumentLoadError("EMPTY_DOCUMENT", f"No readable text was found in {path.name}.")
    return populated


def _load_markdown(path: Path) -> list[SourceSection]:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as error:
        raise DocumentLoadError(
            "CORRUPT_DOCUMENT", f"{path.name} is not valid UTF-8 text."
        ) from error

    sections: list[SourceSection] = []
    heading = "Document"
    content: list[str] = []

    def flush() -> None:
        value = "\n".join(content).strip()
        if value:
            sections.append(
                SourceSection(
                    text=value,
                    locator_type="section",
                    locator_value=heading,
                    section=heading,
                )
            )

    for line in text.splitlines():
        match = re.match(r"^#{1,6}\s+(.+?)\s*#*\s*$", line)
        if match:
            flush()
            heading = match.group(1).strip()
            content = []
        else:
            content.append(line)
    flush()
    return _require_content(sections, path)


def _load_text(path: Path) -> list[SourceSection]:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as error:
        raise DocumentLoadError(
            "CORRUPT_DOCUMENT", f"{path.name} is not valid UTF-8 text."
        ) from error
    return _require_content(
        [
            SourceSection(
                text=text.strip(),
                locator_type="section",
                locator_value="Document",
                section="Document",
            )
        ],
        path,
    )


def _load_pdf(path: Path) -> list[SourceSection]:
    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentLoadError(
                "ENCRYPTED_DOCUMENT", f"{path.name} is password protected."
            )
        sections = [
            SourceSection(
                text=(page.extract_text() or "").strip(),
                locator_type="page",
                locator_value=str(index),
                section=f"Page {index}",
            )
            for index, page in enumerate(reader.pages, start=1)
        ]
    except DocumentLoadError:
        raise
    except Exception as error:
        raise DocumentLoadError(
            "CORRUPT_DOCUMENT", f"{path.name} could not be parsed as a PDF."
        ) from error
    return _require_content(sections, path)


def _load_docx(path: Path) -> list[SourceSection]:
    try:
        document = DocxDocument(str(path))
        sections: list[SourceSection] = []
        heading = "Document"
        content: list[str] = []

        def flush() -> None:
            value = "\n".join(content).strip()
            if value:
                sections.append(
                    SourceSection(
                        text=value,
                        locator_type="section",
                        locator_value=heading,
                        section=heading,
                    )
                )

        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if paragraph.style and paragraph.style.name.startswith("Heading") and value:
                flush()
                heading = value
                content = []
            elif value:
                content.append(value)
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip(" |"):
                    content.append(row_text)
        flush()
    except Exception as error:
        raise DocumentLoadError(
            "CORRUPT_DOCUMENT", f"{path.name} could not be parsed as a DOCX file."
        ) from error
    return _require_content(sections, path)


def _load_pptx(path: Path) -> list[SourceSection]:
    try:
        presentation = Presentation(str(path))
        sections: list[SourceSection] = []
        for index, slide in enumerate(presentation.slides, start=1):
            title_shape = slide.shapes.title
            title = (
                title_shape.text.strip()
                if title_shape and title_shape.text
                else f"Slide {index}"
            )
            values: list[str] = []
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                value = shape.text.strip()
                if value and value != title:
                    values.append(value)
            sections.append(
                SourceSection(
                    text="\n".join(values),
                    locator_type="slide",
                    locator_value=str(index),
                    section=title,
                )
            )
    except Exception as error:
        raise DocumentLoadError(
            "CORRUPT_DOCUMENT", f"{path.name} could not be parsed as a PPTX file."
        ) from error
    return _require_content(sections, path)


LOADERS: dict[str, Loader] = {
    ".md": _load_markdown,
    ".markdown": _load_markdown,
    ".txt": _load_text,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".pptx": _load_pptx,
}


def load_document(path: Path) -> list[SourceSection]:
    """Load supported local content with traceable page, slide, or section metadata."""

    if not path.is_file():
        raise DocumentLoadError("FILE_NOT_FOUND", f"Document {path.name} does not exist.")
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise DocumentLoadError(
            "UNSUPPORTED_EXTENSION",
            f"Files with extension {path.suffix.lower() or '(none)'} are not supported.",
        )
    return loader(path)
