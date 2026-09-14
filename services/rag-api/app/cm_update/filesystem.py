"""Bounded native parsers; no execution, public preview SaaS, or OCR calls."""
from pathlib import Path
import csv
import hashlib
import io
import json
import zipfile
from PIL import Image
from pypdf import PdfReader

MIME={'.pdf':'application/pdf','.txt':'text/plain','.md':'text/plain','.csv':'text/plain',
      '.ipynb':'application/json','.png':'image/png','.jpg':'image/jpeg','.jpeg':'image/jpeg',
      '.webp':'image/webp','.docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
      '.pptx':'application/vnd.openxmlformats-officedocument.presentationml.presentation'}


def parse_document(name: str, content: bytes) -> tuple[str,list[tuple[int,str]],str|None]:
    suffix=Path(name).suffix.lower()
    if suffix not in MIME: raise ValueError('支持 PDF、TXT、Markdown、CSV、Notebook、PNG/JPEG/WebP、DOCX、PPTX')
    chunks: list[tuple[int,str]]=[]
    error=None
    if suffix=='.pdf':
        if not content.startswith(b'%PDF-'): raise ValueError('文件不是有效 PDF')
        try:
            doc=PdfReader(io.BytesIO(content))
            if doc.is_encrypted: raise ValueError('暂不支持加密 PDF')
            if len(doc.pages)>250: raise ValueError('单文件最多 250 页，请拆分上传')
            chunks=[(i+1,(page.extract_text() or '')[:60000]) for i,page in enumerate(doc.pages)]
            if not any(t.strip() for _,t in chunks): error='可预览；未提取到文字。图像题需接入现有 V3 视觉接口。'
        except ValueError: raise
        except Exception: raise ValueError('PDF 无法解析，请检查文件是否损坏') from None
    elif suffix in {'.png','.jpg','.jpeg','.webp'}:
        try:
            with Image.open(io.BytesIO(content)) as im:
                if im.width*im.height>16000000: raise ValueError('图片像素过大')
                actual=im.format
                im.verify()
            expected={'.png':'PNG','.jpg':'JPEG','.jpeg':'JPEG','.webp':'WEBP'}[suffix]
            if actual!=expected: raise ValueError('图片类型与扩展名不一致')
        except Exception: raise ValueError('图片无效或超过像素限制') from None
        error='图片可预览和下载。在学习工作台选择此图片，由已配置的视觉模型读取；不伪造 OCR 结果。'
    elif suffix in {'.docx','.pptx'}:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if len(z.infolist())>4000 or sum(f.file_size for f in z.infolist())>80*1024*1024:
                    raise ValueError('Office 解压大小超限')
            if suffix=='.docx':
                from docx import Document
                doc=Document(io.BytesIO(content))
                parts=[p.text for p in doc.paragraphs]
                parts.extend(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
                chunks=[(1,'\n'.join(parts))]
            else:
                from pptx import Presentation
                doc=Presentation(io.BytesIO(content))
                chunks=[(i+1,'\n'.join(s.text for s in slide.shapes if hasattr(s,'text'))) for i,slide in enumerate(doc.slides)]
        except Exception: raise ValueError('Office 文件解析失败或超过限制') from None
    else:
        try: text=content.decode('utf-8-sig')
        except UnicodeDecodeError: raise ValueError('文本文件请使用 UTF-8 编码') from None
        if '\x00' in text: raise ValueError('文本包含无效二进制字符')
        if suffix=='.ipynb':
            try:
                notebook=json.loads(text)
                text='\n\n'.join(''.join(c.get('source',[])) for c in notebook.get('cells',[]))
            except Exception: raise ValueError('Notebook JSON 无效') from None
        chunks=[(1,text)]
    result=[]
    for page,text in chunks:
        if len(text)>600000: raise ValueError('单页文本过长')
        for start in range(0,len(text),1800):
            part=text[start:start+2100].strip()
            if part: result.append((page,part))
    if len(result)>4000: raise ValueError('文本切分量超限')
    return MIME[suffix],result,error


def valid_filename(name: str) -> str:
    name=Path(name.replace('\\','/')).name.strip()
    if not name or len(name)>220 or any(ord(c)<32 for c in name): raise ValueError('文件名无效')
    return name


def valid_folder(name: str) -> str:
    parts=name.replace('\\','/').strip('/').split('/') if name else []
    if len(name)>250 or any(x in {'.','..',''} or any(ord(c)<32 for c in x) for x in parts):
        raise ValueError('文件夹名称无效')
    return '/'.join(parts)


def file_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()
